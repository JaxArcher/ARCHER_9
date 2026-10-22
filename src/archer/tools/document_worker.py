"""
Worker script executed inside envs/document_tools environment for Phase 5.
Reads JSON request from stdin and outputs JSON response to stdout.
"""

from __future__ import annotations

import io
import json
import os
import sys
from pathlib import Path
from typing import Any, Dict, List


def handle_create_word(data: Dict[str, Any]) -> Dict[str, Any]:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    out_path = Path(data["output_path"])
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    title = data.get("title", "")
    if title:
        p = doc.add_heading(level=0)
        run = p.add_run(title)
        run.font.name = "Calibri"
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)  # Navy accent

    content = data.get("content", "")
    for line in content.split("\n"):
        line_str = line.strip()
        if not line_str:
            continue
        if line_str.startswith("### "):
            doc.add_heading(line_str[4:], level=3)
        elif line_str.startswith("## "):
            doc.add_heading(line_str[3:], level=2)
        elif line_str.startswith("# "):
            doc.add_heading(line_str[2:], level=1)
        elif line_str.startswith("- ") or line_str.startswith("* "):
            doc.add_paragraph(line_str[2:], style="List Bullet")
        elif line_str.startswith("1. ") or (len(line_str) > 3 and line_str[0].isdigit() and line_str[1:3] == ". "):
            idx = line_str.find(". ")
            doc.add_paragraph(line_str[idx + 2:], style="List Number")
        else:
            doc.add_paragraph(line_str)

    doc.save(str(out_path))
    return {
        "status": "success",
        "filepath": str(out_path.resolve()),
        "filename": out_path.name,
        "size_bytes": out_path.stat().st_size,
    }


def handle_edit_word(data: Dict[str, Any]) -> Dict[str, Any]:
    import docx

    in_path = Path(data["filepath"])
    if not in_path.exists():
        return {"status": "error", "error": f"File not found: {in_path}"}

    out_path = Path(data.get("output_path") or in_path)
    doc = docx.Document(str(in_path))

    append_text = data.get("append_text", "")
    for line in append_text.split("\n"):
        line_str = line.strip()
        if not line_str:
            continue
        if line_str.startswith("## "):
            doc.add_heading(line_str[3:], level=2)
        elif line_str.startswith("- ") or line_str.startswith("* "):
            doc.add_paragraph(line_str[2:], style="List Bullet")
        else:
            doc.add_paragraph(line_str)

    doc.save(str(out_path))
    return {
        "status": "success",
        "filepath": str(out_path.resolve()),
        "filename": out_path.name,
        "size_bytes": out_path.stat().st_size,
    }


def handle_create_presentation(data: Dict[str, Any]) -> Dict[str, Any]:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    out_path = Path(data["output_path"])
    out_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = data.get("title", "Presentation")
    if data.get("subtitle"):
        subtitle.text = data.get("subtitle", "")

    # Content slides
    slides_data = data.get("slides", [])
    for s_info in slides_data:
        bullet_slide_layout = prs.slide_layouts[1]
        c_slide = prs.slides.add_slide(bullet_slide_layout)
        c_title = c_slide.shapes.title
        c_title.text = s_info.get("title", "")

        body_shape = c_slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        bullets = s_info.get("bullets", [])
        if bullets:
            tf.text = bullets[0]
            for bullet in bullets[1:]:
                p = tf.add_paragraph()
                p.text = bullet
                p.level = 0

        notes = s_info.get("notes")
        if notes and c_slide.has_notes_slide:
            c_slide.notes_slide.notes_text_frame.text = notes

    prs.save(str(out_path))
    return {
        "status": "success",
        "filepath": str(out_path.resolve()),
        "filename": out_path.name,
        "size_bytes": out_path.stat().st_size,
    }


def handle_create_spreadsheet(data: Dict[str, Any]) -> Dict[str, Any]:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    out_path = Path(data["output_path"])
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    sheets_data = data.get("sheets", [])
    if not sheets_data:
        # Fallback single sheet
        sheets_data = [{
            "name": data.get("title", "Sheet1"),
            "headers": data.get("headers", ["Item", "Value"]),
            "rows": data.get("rows", []),
            "total_columns": data.get("total_columns", []),
        }]

    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    total_font = Font(name="Calibri", size=11, bold=True)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    total_top_border = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

    for s_info in sheets_data:
        sheet_name = s_info.get("name", "Data")[:31]  # Excel 31-char limit
        ws = wb.create_sheet(title=sheet_name)

        headers = s_info.get("headers", [])
        if headers:
            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

        rows = s_info.get("rows", [])
        for row in rows:
            ws.append(row)

        # Style data cells
        for r_idx in range(2, len(rows) + 2):
            for c_idx in range(1, len(headers) + 1):
                ws.cell(row=r_idx, column=c_idx).border = thin_border

        # Add total row if total_columns specified
        total_cols = s_info.get("total_columns", [])
        if total_cols and rows:
            total_row_idx = len(rows) + 2
            ws.cell(row=total_row_idx, column=1, value="Total").font = total_font
            ws.cell(row=total_row_idx, column=1).border = total_top_border

            for col_name in total_cols:
                if col_name in headers:
                    c_idx = headers.index(col_name) + 1
                    col_letter = get_column_letter(c_idx)
                    formula = f"=SUM({col_letter}2:{col_letter}{total_row_idx - 1})"
                    cell = ws.cell(row=total_row_idx, column=c_idx, value=formula)
                    cell.font = total_font
                    cell.border = total_top_border

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(str(out_path))
    return {
        "status": "success",
        "filepath": str(out_path.resolve()),
        "filename": out_path.name,
        "size_bytes": out_path.stat().st_size,
    }


def handle_create_pdf(data: Dict[str, Any]) -> Dict[str, Any]:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    out_path = Path(data["output_path"])
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1B365D"),
        alignment=0,
        spaceAfter=12,
    )
    h1_style = ParagraphStyle(
        "DocH1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#1B365D"),
        spaceBefore=14,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#222222"),
        spaceAfter=6,
    )
    bullet_style = ParagraphStyle(
        "DocBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        leftIndent=15,
        spaceAfter=3,
    )

    story = []
    title = data.get("title", "")
    if title:
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 10))

    content = data.get("content", "")
    for line in content.split("\n"):
        line_str = line.strip()
        if not line_str:
            story.append(Spacer(1, 6))
            continue
        if line_str.startswith("# "):
            story.append(Paragraph(line_str[2:], h1_style))
        elif line_str.startswith("## ") or line_str.startswith("### "):
            story.append(Paragraph(line_str.lstrip("#").strip(), h1_style))
        elif line_str.startswith("- ") or line_str.startswith("* "):
            story.append(Paragraph(f"&bull; {line_str[2:]}", bullet_style))
        else:
            story.append(Paragraph(line_str, body_style))

    doc.build(story)
    return {
        "status": "success",
        "filepath": str(out_path.resolve()),
        "filename": out_path.name,
        "size_bytes": out_path.stat().st_size,
    }


def handle_merge_pdfs(data: Dict[str, Any]) -> Dict[str, Any]:
    from pypdf import PdfWriter

    out_path = Path(data["output_path"])
    out_path.parent.mkdir(parents=True, exist_ok=True)

    pdf_files = data.get("pdf_files", [])
    if not pdf_files:
        return {"status": "error", "error": "No PDF files provided to merge."}

    merger = PdfWriter()
    for f in pdf_files:
        p = Path(f)
        if not p.exists():
            return {"status": "error", "error": f"PDF file not found: {f}"}
        merger.append(str(p.resolve()))

    with open(out_path, "wb") as f_out:
        merger.write(f_out)
    merger.close()

    return {
        "status": "success",
        "filepath": str(out_path.resolve()),
        "filename": out_path.name,
        "size_bytes": out_path.stat().st_size,
    }


def main():
    try:
        raw_in = sys.stdin.read()
        if not raw_in.strip():
            print(json.dumps({"status": "error", "error": "Empty input payload"}))
            return

        payload = json.loads(raw_in)
        action = payload.get("action")

        handlers = {
            "create_word": handle_create_word,
            "edit_word": handle_edit_word,
            "create_presentation": handle_create_presentation,
            "create_spreadsheet": handle_create_spreadsheet,
            "create_pdf": handle_create_pdf,
            "merge_pdfs": handle_merge_pdfs,
        }

        handler = handlers.get(action)
        if not handler:
            print(json.dumps({"status": "error", "error": f"Unknown action: {action}"}))
            return

        result = handler(payload.get("data", {}))
        print(json.dumps(result))
    except Exception as e:
        import traceback
        print(json.dumps({"status": "error", "error": str(e), "traceback": traceback.format_exc()}))


if __name__ == "__main__":
    main()
